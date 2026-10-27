#!/usr/bin/env python3
"""The Information Sheet's blocks, read by loto-web's own reader.

Since the more-than-10-sources change
(loto-web/docs/Field_App_More_Than_10_Sources_Spec_2026-10-27.md, F5) a unit's
block on the sheet grows past 10 sources, and loto-web finds each block with
infosheet_parser.grid_blocks (anchored at its "Equipment ID/Name" header and
measured by its "Tied to" row). This script opens export ZIPs from the app,
runs loto-web's reader on every Information Sheet in them, and checks each
block against its unit in entries.json:

  - the block lines up, with no label warning, and has max(10, sources) slots;
  - its sources, slot by slot: energy source, device, device ID, quantity,
    location and photo;
  - Tied to, Template and Notes, and the unit answers (Shop, Change Later,
    UV Lights, ...) where loto-web reads them, moved down with the Tied to row.

Make an export in the app (any export; a page can POST one to
tests/fuzz/fuzz_receiver.py), then:
    LOTO_WEB=../loto-web python3 tests/contract/check_sheet_blocks.py export.zip [more.zip ...]
(LOTO_WEB defaults to the loto-web folder beside this repo's folder.)
Exit 0 when every block agrees. It only reads loto-web; it writes nothing there.
"""
import io
import json
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
LOTO_WEB = os.path.abspath(os.environ.get("LOTO_WEB") or os.path.join(os.path.dirname(REPO), "loto-web"))
sys.path.insert(0, LOTO_WEB)

import openpyxl  # noqa: E402
from app.services import infosheet_parser as isp  # noqa: E402

# entries.json key -> the sheet label it travels under (the app's unitAnswersForJson)
ANSWER_KEYS = {
    "Shop": lambda e: e.get("shop"),
    "Change Later": lambda e: (e.get("changeLater") or {}).get("text"),
    "UV Lights": lambda e: e.get("uvLights"),
    "Confined Space": lambda e: e.get("confinedSpace"),
    "ATS Make": lambda e: e.get("atsMake"),
    "Backfeed": lambda e: "; ".join(e.get("atsBackfeed") or []),
    "Capacitors": lambda e: e.get("capacitors"),
    "Exhaust": lambda e: e.get("vacuumExhaust"),
    "Union/Flange": lambda e: e.get("reliefJoint"),
    "Refrigerant": lambda e: (e.get("refrigerant") or {}).get("type"),
    "Refrigerant PSI": lambda e: (e.get("refrigerant") or {}).get("psi"),
    "Damper Actuators": lambda e: "None" if e.get("dampersNone") else None,
    "Alternate Sources": lambda e: e.get("alternateSources"),
    "Terminal Access": lambda e: e.get("terminalAccess"),
    "Transformer Nameplate": lambda e: e.get("transformerNameplate"),
    "SPD / Capacitors / Grounding": lambda e: e.get("spdCapacitorsGrounding"),
}


def text(v):
    return "" if v is None else " ".join(str(v).split())


def check(zpath, problems, counts):
    z = zipfile.ZipFile(zpath)
    entries = json.loads(z.read("entries.json"))["entries"]
    unread = {e["id"] for e in entries}
    sheets = sorted(n for n in z.namelist() if re.match(r"info_sheets/Information_Sheet_.*\.xlsx$", n))
    for sp in sheets:
        where = f"{os.path.basename(zpath)} {os.path.basename(sp)}"
        ws = openpyxl.load_workbook(io.BytesIO(z.read(sp)))["InfoSheetForm"]
        layout = isp._detect_info_sheet_layout(ws)
        if not layout or layout["name"] != isp.LAYOUT_FIELDAPP["name"]:
            problems.append(f"{where}: not read as the field-app layout")
            continue
        blocks = {b.base_row: b for b in isp.read_grid_blocks(ws, layout, sp)}
        for g in isp.grid_blocks(ws, layout):
            name, place = text(ws.cell(g.base_row + 3, 1).value), text(ws.cell(g.base_row + 3, 4).value)
            tag = f'{where} "{name}" (header row {g.header_row})'
            counts["blocks"] += 1
            if g.problem:
                problems.append(f"{tag}: loto-web would skip it: {g.problem}")
                continue
            if g.label_note:
                problems.append(f"{tag}: label warning: {g.label_note}")
            mine = [e for e in entries if e["id"] in unread and text(e.get("equipName")) == name
                    and text(" ".join(filter(None, [e.get("equipBuilding"), e.get("equipRoom")]))) == place]
            if len(mine) != 1:
                problems.append(f"{tag}: {len(mine)} units in entries.json with that name and location")
                continue
            e = mine[0]
            unread.discard(e["id"])
            srcs = e.get("sources") or []
            if g.slots != max(isp.MAX_SOURCES_PER_PAGE, len(srcs)):
                problems.append(f"{tag}: {g.slots} source slots for {len(srcs)} sources")
            b = blocks.get(g.base_row)
            if b is None:
                problems.append(f"{tag}: read_grid_blocks returned no block")
                continue
            if len(b.sources) != len(srcs):
                problems.append(f"{tag}: loto-web reads {len(b.sources)} sources, entries.json has {len(srcs)}")
            for i, (s, r) in enumerate(zip(srcs, b.sources), 1):
                counts["sources"] += 1
                want = {"slot": i, "source": text(s.get("energySource")) or "Unknown", "device": text(s.get("deviceType")),
                        "device_id": text(s.get("deviceId")), "qty": int(s.get("quantity") or 1),
                        "location": text(s.get("location")), "photo": os.path.basename(s.get("photoFile") or "")}
                have = {"slot": r.slot, "source": text(r.raw_name), "device": text(r.device), "device_id": text(r.device_id),
                        "qty": r.qty, "location": text(r.location), "photo": text(r.photo_ref)}
                for k in want:
                    if want[k] != have[k]:
                        problems.append(f"{tag} source {i}: {k} {want[k]!r} in entries.json, {have[k]!r} on the sheet")
            tied = text(e.get("tiedTo"))
            if (tied and not text(b.tied_to).startswith(tied)) or (not tied and b.tied_to):
                problems.append(f"{tag}: Tied to {b.tied_to!r}, entries.json {tied!r}")
            for k, have in (("template", b.template), ("notes", b.notes)):
                if text(e.get(k)) != text(have):
                    problems.append(f"{tag}: {k} {have!r} on the sheet, {e.get(k)!r} in entries.json")
            answers = isp.read_unit_answers(ws, g.base_row, layout, g.shift)
            for label in ("Shop", "Change Later"):
                if label not in answers:
                    problems.append(f"{tag}: no {label!r} label where loto-web looks (summary row {g.summary_row})")
            for label, value in answers.items():
                counts["answers"] += 1
                want = text(ANSWER_KEYS[label](e)) if label in ANSWER_KEYS else None
                if want is None or want != text(value):
                    problems.append(f"{tag}: answer {label!r} = {value!r} on the sheet, {want!r} in entries.json")
    for e in entries:
        if e["id"] in unread:
            problems.append(f'{os.path.basename(zpath)}: "{e.get("equipName")}" has no block loto-web reads')


def main():
    zips = sys.argv[1:]
    if not zips:
        sys.exit(__doc__)
    problems, counts = [], {"blocks": 0, "sources": 0, "answers": 0}
    for zp in zips:
        check(zp, problems, counts)
    print(json.dumps({"zips": len(zips), "counts": counts, "problems": problems[:40], "n_problems": len(problems)}, indent=1))
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
